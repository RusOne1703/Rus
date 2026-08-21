import docx

doc = docx.Document()
doc.add_heading('Где провести семинар для организации в СПб: Топ-15 лучших площадок (2025)', 0)

with open('article_dzen_spb_seminars.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

for line in lines:
    line_str = line.strip()
    if not line_str:
        continue
    if line_str.startswith('# '):
        doc.add_heading(line_str[2:], level=1)
    elif line_str.startswith('## '):
        doc.add_heading(line_str[3:], level=2)
    elif line_str.startswith('### '):
        doc.add_heading(line_str[4:], level=3)
    elif line_str.startswith('* '):
        doc.add_paragraph(line_str[2:], style='List Bullet')
    elif line_str.startswith('- '):
        doc.add_paragraph(line_str[2:], style='List Bullet')
    else:
        # Clean markdown bold/italic tags for clean text reading
        clean_text = line_str.replace('**', '').replace('*', '').replace('`', '').replace('>', '')
        doc.add_paragraph(clean_text)

doc.save('article_dzen_spb_seminars.docx')
print("Successfully generated article_dzen_spb_seminars.docx")
