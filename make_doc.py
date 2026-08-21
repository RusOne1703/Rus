import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Page Setup - Margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Color Palette
COLOR_PRIMARY = RGBColor(26, 54, 93)      # Dark Blue #1A365D
COLOR_SECONDARY = RGBColor(197, 48, 48)   # Crimson Accent #C53030
COLOR_TEXT = RGBColor(45, 55, 72)         # Dark Gray #2D3748
COLOR_MUTED = RGBColor(113, 128, 150)     # Muted Slate #718096
COLOR_PROMPT = RGBColor(44, 122, 123)     # Teal #2C7A7B

# Base Style setup
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)
style_normal.font.color.rgb = COLOR_TEXT

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_callout_box(doc, text, title=None, fill_hex="F7FAFC", border_hex="3182CE"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    set_cell_background(cell, fill_hex)

    # Border
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.right_indent = Inches(0.1)

    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = COLOR_PRIMARY

    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)
    run_text.font.italic = True
    run_text.font.color.rgb = COLOR_TEXT

    # Spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

print("Helper functions ready")
