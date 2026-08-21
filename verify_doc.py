import docx

doc = docx.Document('топ_15_площадок_для_семинара_спб.docx')

urls = ["luminis-loft.ru", "memoriesloft.ru", "eventpandora.ru", "show-ring.ru"]

all_text = ""
for p in doc.paragraphs:
    all_text += p.text + "\n"
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            all_text += cell.text + "\n"

print("--- VERIFICATION CHECK ---")

for url in urls:
    count = all_text.lower().count(url.lower())
    print(f"URL '{url}': found {count} time(s)")
    assert count >= 1, f"Missing URL: {url}"

# Check for keywords
keywords = ["семинар", "спб", "санкт-петербург", "площадк", "лофт", "стоимость", "адрес", "плюсы", "минусы", "для кого"]
for kw in keywords:
    count = all_text.lower().count(kw)
    print(f"Keyword '{kw}': found {count} time(s)")
    assert count >= 1, f"Missing keyword: {kw}"

# Check number of venues (1. to 15.)
for i in range(1, 16):
    assert f"{i}." in all_text, f"Missing venue number {i}"

# Check table count (1 main prompt + 15 venue prompts)
print(f"Total Callout Tables (Prompts): {len(doc.tables)}")
assert len(doc.tables) == 16, f"Expected 16 tables, got {len(doc.tables)}"

print("\n✅ ALL VERIFICATION CHECKS PASSED SUCCESSFULLY!")
